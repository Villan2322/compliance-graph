"""Secure Controls Framework workbook -> (:SCFControl) and cross-framework
(:Requirement)-[:MAPS_TO]->(:SCFControl) edges. SCF is the hub that lets one
finding light up NIST, ISO, SOC 2, PCI, HIPAA and more at once.

For Tier 3 frameworks (ISO, SOC 2, PCI, CIS) only the native ID is stored;
the requirement text is never written to the graph from this loader."""
from __future__ import annotations

import re
from pathlib import Path

import yaml

from ..config import ROOT
from ..ids import nist_80053, req_id

COLUMNS = ROOT / "sources" / "scf_columns.yaml"
_SPLIT = re.compile(r"[\n\r;,]+")


def _cfg() -> dict:
    return yaml.safe_load(COLUMNS.read_text())


def _find(headers: list[str], pattern: str) -> int | None:
    rx = re.compile(pattern, re.I)
    for i, h in enumerate(headers):
        if h and rx.search(str(h).strip()):
            return i
    return None


def headers(path: Path) -> dict[str, list[str]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        out[ws.title] = [str(h) for h in first if h]
    return out


def _native(fw: str, token: str) -> str | None:
    token = token.strip()
    if not token or token.lower() in {"n/a", "none", "-"}:
        return None
    if fw == "NIST-800-53-r5":
        return nist_80053(token)
    return token


def parse(path: Path) -> dict:
    from openpyxl import load_workbook
    cfg = _cfg()
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = next((w for w in wb.worksheets if re.search(cfg["sheet"], w.title, re.I)), wb.worksheets[0])
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h else "" for h in next(rows)]
    col = {k: _find(hdr, cfg[k]) for k in ("control_id", "control_title", "control_text", "domain", "weight")}
    if col["control_id"] is None:
        raise ValueError(f"SCF control id column not found in sheet {ws.title!r}; run `cgraph scf-headers`")
    fw_cols = {fw: _find(hdr, pat) for fw, pat in cfg["frameworks"].items()}
    missing = [fw for fw, i in fw_cols.items() if i is None]
    controls, maps = [], []
    for r in rows:
        sid = r[col["control_id"]]
        if not sid:
            continue
        sid = str(sid).strip()
        get = lambda k: (str(r[col[k]]).strip() if col[k] is not None and r[col[k]] is not None else "")
        controls.append({"scfId": sid, "title": get("control_title"), "text": get("control_text"),
                         "domain": get("domain"), "weight": get("weight")})
        for fw, i in fw_cols.items():
            if i is None or r[i] is None:
                continue
            for tok in _SPLIT.split(str(r[i])):
                n = _native(fw, tok)
                if n:
                    maps.append({"fw": fw, "native": n, "reqId": req_id(fw, n), "scf": sid})
    return {"controls": controls, "maps": maps, "sheet": ws.title, "missingFrameworks": missing}


def load(graph, data: dict, spec) -> dict:
    sid = spec.id
    graph.batch("""
        UNWIND $rows AS r MERGE (c:SCFControl {scfId: r.scfId})
        SET c.title = r.title, c.text = r.text, c.domain = r.domain, c.weight = r.weight,
            c.sourceId = $sid, c.ingestedAt = datetime(), c.isActive = true
        WITH c MATCH (f:Framework {frameworkId: 'SCF'}) MERGE (c)-[p:PART_OF]->(f) SET p.sourceId = $sid""",
        data["controls"], sid=sid)
    n = graph.batch("""
        UNWIND $rows AS r
        MATCH (f:Framework {frameworkId: r.fw}), (c:SCFControl {scfId: r.scf})
        MERGE (q:Requirement {reqId: r.reqId})
        ON CREATE SET q.frameworkId = r.fw, q.nativeId = r.native, q.isActive = true,
                      q.textLicensed = f.textRedistributable, q.sourceId = $sid, q.ingestedAt = datetime()
        MERGE (q)-[p:PART_OF]->(f) ON CREATE SET p.sourceId = $sid
        MERGE (q)-[m:MAPS_TO]->(c)
        SET m.sourceId = $sid, m.method = 'scf-crosswalk', m.strm = null, m.confidence = 0.8""",
        data["maps"], sid=sid)
    return {"controls": len(data["controls"]), "crosswalkEdges": n,
            "sheet": data["sheet"], "missingFrameworks": ",".join(data["missingFrameworks"])}
