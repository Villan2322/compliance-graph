"""Risk Control Matrix outputs: XLSX + CSV for people, OSCAL-shaped JSON for GRC
tools, Markdown for PR comments."""
from __future__ import annotations

import csv
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path

COLUMNS = [
    "Risk ID", "Rating", "Score", "Likelihood", "Impact", "Risk Statement", "CWE", "Weakness",
    "ATT&CK Techniques", "ATLAS Techniques", "CVEs", "Control ID (NIST 800-53 r5)", "Control Title",
    "Mapping Path", "Mapping Confidence", "Mapping Reviewed", "SCF Control(s)", "Crosswalk (other frameworks)",
    "Control Type", "Control Owner", "Test Procedure", "Evidence Required", "Affected Files", "Detected By",
    "Decision", "Status", "Scoring Rationale",
]

DETECTIVE = ("RA-5", "SI-4", "AU-", "SA-11", "SA-15", "CA-")


def control_type(native: str) -> str:
    return "Detective" if native.startswith(DETECTIVE) else "Preventive"


def rows(risks: list[dict]) -> list[dict]:
    out = []
    for r in risks:
        controls = r["controls"] or [{"nativeId": "UNMAPPED", "title": "No control mapped — needs analyst review",
                                      "path": "none", "reqId": None}]
        for c in controls:
            cw = r["crosswalk"].get(c.get("reqId") or "", {})
            scf_ids = [s for s in cw.get("scf", []) if s]
            mapped = "; ".join(sorted({f"{m['frameworkId']} {m['nativeId']}" for m in cw.get("mapped", [])}))[:1500]
            conf = c.get("confidence")
            out.append({
                "Risk ID": r["riskId"], "Rating": r["rating"], "Score": r["score"],
                "Likelihood": r["likelihood"], "Impact": r["impact"], "Risk Statement": r["statement"],
                "CWE": r["cwe"], "Weakness": r["cweName"],
                "ATT&CK Techniques": ", ".join(r["techniques"]), "ATLAS Techniques": ", ".join(r["aiTechniques"]),
                "CVEs": ", ".join(r["cves"]),
                "Control ID (NIST 800-53 r5)": c["nativeId"], "Control Title": c.get("title", ""),
                "Mapping Path": {"direct": "Curated CWE→control", "threat": "CWE→CAPEC→ATT&CK→control (CTID)"}.get(c.get("path"), "—"),
                "Mapping Confidence": f"{conf:.2f}" if isinstance(conf, (int, float)) else (f"{c.get('hits')} technique path(s)" if c.get("hits") else ""),
                "Mapping Reviewed": "Yes" if c.get("reviewed") else ("No — flagged for review" if c.get("path") == "direct" else "Upstream (CTID)"),
                "SCF Control(s)": ", ".join(s["scfId"] for s in scf_ids),
                "Crosswalk (other frameworks)": mapped or ("No crosswalk mapped in SCF for this control" if scf_ids else "Load SCF to populate"),
                "Control Type": control_type(c["nativeId"]), "Control Owner": r["owner"],
                "Test Procedure": (f"Re-run {', '.join(r['tools'])} ({', '.join(r['rules'][:3])}) on {', '.join(r['files'][:3]) or 'the component'}; "
                                   f"confirm zero findings for {r['cwe']}. Inspect the fix against {c['nativeId']} "
                                   f"({c.get('title', '')}) and confirm the control is enforced in code review."),
                "Evidence Required": ("CI scan artifact showing the rule passing; merged PR/commit with the fix; "
                                      "code-review record; control-owner attestation."),
                "Affected Files": ", ".join(r["files"][:10]), "Detected By": ", ".join(r["tools"]),
                "Decision": r.get("decision") or "", "Status": r["status"],
                "Scoring Rationale": "; ".join(r["scoringRationale"]),
            })
    return out


def write_csv(risks: list[dict], path: Path) -> Path:
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows(risks))
    return path


def write_xlsx(risks: list[dict], path: Path, meta: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Risk Control Matrix"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    fills = {"Critical": "FECACA", "High": "FED7AA", "Medium": "FEF08A", "Low": "D9F99D"}
    for r in rows(risks):
        ws.append([r[k] for k in COLUMNS])
        ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=fills.get(r["Rating"], "FFFFFF"))
    widths = {"Risk Statement": 60, "Test Procedure": 60, "Crosswalk (other frameworks)": 45, "Scoring Rationale": 50,
              "Control Title": 35, "Evidence Required": 40, "Affected Files": 35}
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 16)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet("Run")
    for k, v in meta.items():
        info.append([k, v if isinstance(v, (str, int, float)) else json.dumps(v)])
    wb.save(path)
    return path


def write_oscal(risks: list[dict], path: Path, meta: dict) -> Path:
    """OSCAL 1.1.x assessment-results *shaped* JSON. Validate with oscal-cli
    before handing to a GRC platform; see docs/ROADMAP.md."""
    now = datetime.now(timezone.utc).isoformat()
    observations, oscal_risks, findings, controls = [], [], [], set()
    for r in risks:
        obs_id, risk_uuid = str(uuid.uuid4()), str(uuid.uuid4())
        observations.append({"uuid": obs_id, "title": f"{r['cwe']} detected by {', '.join(r['tools'])}",
                             "description": f"{len(r['findingIds'])} finding(s) in {', '.join(r['files'][:5])}",
                             "methods": ["TEST"], "types": ["finding"], "collected": now})
        oscal_risks.append({"uuid": risk_uuid, "title": f"{r['riskId']} {r['cweName']}", "description": r["statement"],
                            "statement": "; ".join(r["scoringRationale"]), "status": "open",
                            "characterizations": [{"origin": {"actors": [{"type": "tool", "actor-uuid": meta["tool_uuid"]}]},
                                                   "facets": [{"name": "likelihood", "system": "https://github.com/compliance-graph", "value": str(r["likelihood"])},
                                                              {"name": "impact", "system": "https://github.com/compliance-graph", "value": str(r["impact"])}]}],
                            "related-observations": [{"observation-uuid": obs_id}]})
        for c in r["controls"]:
            if c.get("oscalId"):
                controls.add(c["oscalId"])
                findings.append({"uuid": str(uuid.uuid4()), "title": f"{c['nativeId']} not satisfied for {r['cwe']}",
                                 "description": r["statement"],
                                 "target": {"type": "objective-id", "target-id": f"{c['oscalId']}_obj",
                                            "status": {"state": "not-satisfied"}},
                                 "related-observations": [{"observation-uuid": obs_id}],
                                 "related-risks": [{"risk-uuid": risk_uuid}]})
    doc = {"assessment-results": {
        "uuid": str(uuid.uuid4()),
        "metadata": {"title": f"compliance-graph audit {meta['run_id']}", "last-modified": now,
                     "version": meta["run_id"], "oscal-version": "1.1.2"},
        "import-ap": {"href": "#assessment-plan-not-provided"},
        "results": [{"uuid": str(uuid.uuid4()), "title": f"Automated code audit of {meta['repo_id']}",
                     "description": "Generated by compliance-graph audit loop.", "start": meta["started_at"], "end": now,
                     "reviewed-controls": {"control-selections": [{"include-controls": [{"control-id": c} for c in sorted(controls)]}]},
                     "observations": observations, "risks": oscal_risks, "findings": findings}]}}
    path.write_text(json.dumps(doc, indent=2))
    return path


def write_markdown(risks: list[dict], path: Path, meta: dict, summary: str | None) -> Path:
    L = [f"# Audit {meta['run_id']} — {meta['repo_id']}", ""]
    if summary:
        L += [summary.strip(), ""]
    counts = {k: sum(1 for r in risks if r["rating"] == k) for k in ("Critical", "High", "Medium", "Low")}
    L += ["**" + " · ".join(f"{k}: {v}" for k, v in counts.items()) + "**", "",
          "| Risk | Rating | CWE | Top controls | Decision |", "|---|---|---|---|---|"]
    for r in risks:
        ctl = ", ".join(c["nativeId"] for c in r["controls"][:3]) or "unmapped"
        L.append(f"| {r['riskId']} | {r['rating']} ({r['score']}) | {r['cwe']} {r['cweName']} | {ctl} | {r.get('decision') or '—'} |")
    if meta.get("warnings"):
        L += ["", "**Warnings**", *[f"- {w}" for w in meta["warnings"]]]
    L += ["", "_Unreviewed curated mappings are flagged in the RCM. Treat this as decision support, not an attestation._"]
    path.write_text("\n".join(L) + "\n")
    return path
